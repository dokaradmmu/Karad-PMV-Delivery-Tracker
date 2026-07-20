"""
Excel output writer for the PMV Toolkit Delivery Tracker (Karad v1).

Kept separate from pmv_core.py so the processing logic and the file-format
concerns don't tangle. No frozen panes (per user instruction). No formulas
are needed here — this is a data listing/reporting workbook, not a
financial model — so there is nothing to recalculate.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from pmv_core import ProcessResult

FONT_NAME = "Arial"
HEADER_FILL_COLOR = "1F3864"  # matches the navy used in the user's other reports


def _write_sheet(ws: Worksheet, df: pd.DataFrame, title_note: str | None = None) -> None:
    start_row = 1

    if title_note:
        ws.cell(row=1, column=1, value=title_note)
        ws.cell(row=1, column=1).font = Font(name=FONT_NAME, size=12, bold=True, color="1F3864")
        start_row = 3  # leave a blank row after the title

    header_row = start_row
    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=col_name)
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
        cell.fill = __import__("openpyxl").styles.PatternFill(
            start_color=HEADER_FILL_COLOR, end_color=HEADER_FILL_COLOR, fill_type="solid"
        )
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for row_idx, row in enumerate(df.itertuples(index=False), start=header_row + 1):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = Font(name=FONT_NAME, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=False)

    # Reasonable column widths based on header + a sample of content length.
    for col_idx, col_name in enumerate(df.columns, start=1):
        sample_lengths = [len(str(col_name))]
        if len(df) > 0:
            sample_lengths += [len(str(v)) for v in df[col_name].astype(str).head(50)]
        width = min(max(sample_lengths) + 2, 45)
        ws.column_dimensions[get_column_letter(col_idx)].width = max(width, 10)

    ws.row_dimensions[header_row].height = 28


def _safe_sheet_title(title: str) -> str:
    """Excel sheet names: max 31 chars, no []:*?/\\ characters."""
    for ch in "[]:*?/\\":
        title = title.replace(ch, "-")
    return title[:31]


def write_workbook(result: ProcessResult, output_path):
    """
    `output_path` may be a filesystem path (str/Path) or a writable
    file-like object (e.g. io.BytesIO), the latter used by the Streamlit
    app to avoid writing to a shared on-disk path across concurrent users.
    """
    if isinstance(output_path, (str, Path)):
        output_path = Path(output_path)

    wb = Workbook()
    wb.remove(wb.active)

    sheets: list[tuple[str, pd.DataFrame, str | None]] = [
        ("Delivered", result.delivered, None),
        ("Pending", result.pending, None),
        ("Returned to Sender", result.returned_to_sender, None),
        (
            f"Delivered {result.old_date_label}-{result.new_date_label}",
            result.newly_delivered,
            f"Toolkits delivered between {result.old_date_label} and {result.new_date_label}",
        ),
        ("Other Divisions", result.other_divisions, None),
    ]

    for sheet_title, df, note in sheets:
        ws = wb.create_sheet(_safe_sheet_title(sheet_title))
        _write_sheet(ws, df, title_note=note)

    wb.save(output_path)
    return output_path
